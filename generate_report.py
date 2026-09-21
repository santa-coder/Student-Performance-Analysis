from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

TEMPLATE_FILE = BASE_DIR / "student report format.docx"
OUTPUT_FILE = BASE_DIR / "Student_Learning_Pattern_Analysis_Project_Report.docx"


PROJECT_TITLE = "STUDENT LEARNING PATTERN ANALYSIS USING CLUSTERING"

STUDENT_1 = "C SRIKISHORE"
REG_1 = "610524091079"

STUDENT_2 = "RITHIK VISHNU"
REG_2 = "610524091059"

STUDENT_3 = "D SABARIVEL"
REG_3 = "610524091060"

DEPARTMENT = "DEPARTMENT OF ARTIFICIAL INTELLIGENCE AND DATA SCIENCE"
COLLEGE = "HINDUSTHAN COLLEGE OF TECHNOLOGY"
LOCATION = "SALEM – 636309"


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def set_font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold


def add_paragraph(doc, text="", bold=False, size=12, align=None):
    p = doc.add_paragraph()

    if align is not None:
        p.alignment = align

    p.paragraph_format.line_spacing = 1.5

    run = p.add_run(text)
    set_font(run, size, bold)

    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()

    if level == 1:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)

        run = p.add_run(text)
        set_font(run, 14, True)

    elif level == 2:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)

        run = p.add_run(text)
        set_font(run, 12, True)

    else:
        run = p.add_run(text)
        set_font(run, 12, True)

    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()

    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-10)
    p.paragraph_format.line_spacing = 1.5

    run = p.add_run("• " + text)
    set_font(run, 12)

    return p

def add_number(doc, text):
    # Simple numbered paragraph without relying on template styles
    if not hasattr(add_number, "counter"):
        add_number.counter = 0

    add_number.counter += 1

    p = doc.add_paragraph()

    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-18)
    p.paragraph_format.line_spacing = 1.5

    run = p.add_run(f"{add_number.counter}. {text}")
    set_font(run, 12)

    return p

def add_figure_placeholder(doc, number, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(
        f"\n\n[ INSERT PROJECT SCREENSHOT / FIGURE HERE ]\n\n"
    )
    set_font(run, 11, False)

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = caption.add_run(f"Figure {number}: {title}")
    set_font(run, 11, True)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header

        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_font(run, 10, True)

    for row in rows:
        cells = table.add_row().cells

        for i, value in enumerate(row):
            cells[i].text = str(value)

            for paragraph in cells[i].paragraphs:
                for run in paragraph.runs:
                    set_font(run, 10)

    return table


def replace_text_in_paragraph(paragraph, old, new):
    if old in paragraph.text:
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)


def replace_everywhere(doc, old, new):

    # Normal paragraphs
    for paragraph in doc.paragraphs:
        replace_text_in_paragraph(paragraph, old, new)

    # Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_text_in_paragraph(paragraph, old, new)

    # Headers and footers
    for section in doc.sections:

        for paragraph in section.header.paragraphs:
            replace_text_in_paragraph(paragraph, old, new)

        for paragraph in section.footer.paragraphs:
            replace_text_in_paragraph(paragraph, old, new)


def set_document_font(doc):

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"


# ============================================================
# LOAD TEMPLATE
# ============================================================

if not TEMPLATE_FILE.exists():
    print("ERROR: Template file not found.")
    print()
    print("Expected:")
    print(TEMPLATE_FILE)
    print()
    print("Make sure 'student report format.docx' is in the same folder.")
    raise SystemExit(1)


print("Loading template...")

doc = Document(TEMPLATE_FILE)


# ============================================================
# REPLACE BASIC TEMPLATE PLACEHOLDERS
# ============================================================

replacements = {

    "[PROJECT TITLE]": PROJECT_TITLE,
    "[STUDENT NAME 1]": STUDENT_1,
    "[REGISTER NUMBER 1]": REG_1,
    "[STUDENT NAME 2]": STUDENT_2,
    "[REGISTER NUMBER 2]": REG_2,
    "[STUDENT NAME 3]": STUDENT_3,
    "[REGISTER NUMBER 3]": REG_3,

    "[DEPARTMENT NAME]": DEPARTMENT,
    "[COLLEGE NAME]": COLLEGE,
    "[LOCATION]": LOCATION,

    "September 2026": "September 2026",
}


for old, new in replacements.items():
    replace_everywhere(doc, old, new)


# ============================================================
# ADD PROJECT CONTENT
# ============================================================

# ------------------------------------------------------------
# ABSTRACT
# ------------------------------------------------------------

add_heading(doc, "ABSTRACT", 1)

add_paragraph(
    doc,
    """Student performance can vary based on attendance, academic scores and other measurable learning indicators. Identifying common patterns can help teachers understand groups of students and provide suitable academic support. This project presents a simple Student Learning Pattern Analysis system using K-Means clustering to group students according to similarity in selected academic and learning-related attributes."""
)

add_paragraph(
    doc,
    """The system accepts structured student data, validates the records and preprocesses the numerical features before clustering. Standardization is applied so that features measured on different scales can be used together. K-Means then assigns students to clusters based on their distance from cluster centroids. The resulting groups are displayed through an interactive dashboard containing performance summaries and graphical analysis."""
)

add_paragraph(
    doc,
    """The application is implemented using Python, Streamlit, SQLite, Pandas, Scikit-learn and Plotly. SQLite provides persistent storage, while the Streamlit interface supports role-based access for students, teachers and administrators. The project also provides individual student analysis, intervention records and Excel report generation."""
)

add_paragraph(
    doc,
    """The system is intended as a decision-support tool for academic monitoring. It does not replace teacher judgement or automatically determine a student's ability. The final application demonstrates how a basic machine-learning technique can be combined with data visualization and database management to make student learning patterns easier to understand and support timely academic intervention."""
)

add_paragraph(
    doc,
    "Keywords: Student Performance, K-Means Clustering, Machine Learning, Learning Patterns, Data Visualization"
)


# ============================================================
# CHAPTER 1
# ============================================================

add_heading(doc, "CHAPTER 1", 1)
add_heading(doc, "INTRODUCTION", 1)

add_heading(doc, "1.1 Background", 2)

add_paragraph(
    doc,
    """Educational institutions collect different types of student information such as attendance, academic marks, grade point average and other learning indicators. Analysing these records manually can make it difficult to identify common learning patterns, especially when the number of students increases."""
)

add_paragraph(
    doc,
    """Machine learning techniques can be used to discover useful patterns from student data. Clustering is an unsupervised learning technique that groups similar data records together. In this project, K-Means clustering is used to identify groups of students with similar academic and learning characteristics."""
)

add_heading(doc, "1.2 Problem Statement", 2)

add_paragraph(
    doc,
    """Traditional student monitoring commonly depends on spreadsheets, individual records and manual interpretation. Although these methods can work for small datasets, they require repeated checking and may not clearly show groups of students with similar performance patterns."""
)

add_paragraph(
    doc,
    """Therefore, a simple system is required to preprocess student data, apply clustering and display the resulting learning patterns through an understandable dashboard."""
)

add_heading(doc, "1.3 Motivation", 2)

add_paragraph(
    doc,
    """The main motivation of this project is to provide teachers and academic staff with a quick and understandable method for analysing student learning patterns. Grouping students based on similar characteristics can support mentoring, academic follow-up and remedial activities."""
)

add_heading(doc, "1.4 Objectives", 2)

objectives = [
    "To collect and store structured student academic information.",
    "To validate and preprocess the input student data.",
    "To apply K-Means clustering to identify similar student groups.",
    "To visualize clustering and performance information through an interactive dashboard.",
    "To provide individual student performance analysis.",
    "To maintain intervention and teacher feedback records.",
    "To generate student analysis reports in Excel format."
]

for item in objectives:
    add_bullet(doc, item)


add_heading(doc, "1.5 Scope of the Project", 2)

add_paragraph(
    doc,
    """The project focuses on analysing structured student performance data using clustering. It includes data validation, preprocessing, feature selection, standardization, K-Means clustering, graphical visualization and student-level analysis. The application also provides role-based access for students, teachers and administrators."""
)

add_heading(doc, "1.6 Applications", 2)

applications = [
    "Student performance grouping.",
    "Academic mentoring and follow-up.",
    "Attendance-based monitoring.",
    "Identification of students requiring additional support.",
    "Performance trend analysis.",
    "Generation of student reports."
]

for item in applications:
    add_bullet(doc, item)


add_heading(doc, "1.7 Project Contributions", 2)

add_paragraph(
    doc,
    """The project integrates machine learning, data visualization, database management and reporting into a single academic monitoring application. K-Means clustering is combined with Streamlit for the user interface, SQLite for persistent data storage, Plotly for visualization and Excel generation for reports."""
)

add_heading(doc, "1.8 Organization of the Report", 2)

add_paragraph(
    doc,
    """Chapter 1 introduces the project. Chapter 2 presents the literature survey. Chapter 3 explains the existing system. Chapter 4 describes the proposed system and algorithm. Chapter 5 presents system requirements. Chapter 6 explains methodology and implementation. Chapter 7 presents results and discussion. Chapter 8 concludes the project and discusses future scope.""")


# ============================================================
# CHAPTER 2
# ============================================================

add_heading(doc, "CHAPTER 2", 1)
add_heading(doc, "LITERATURE SURVEY", 1)

add_heading(doc, "2.1 Introduction", 2)

add_paragraph(
    doc,
    """The literature survey provides background for clustering, data preprocessing and educational data analysis. Previous work on clustering has established methods for grouping numerical observations based on similarity. These concepts form the foundation of the proposed student learning pattern analysis system."""
)

add_heading(doc, "2.2 Review of Existing Research", 2)

add_paragraph(
    doc,
    """K-Means is one of the commonly used clustering techniques for numerical datasets. It partitions observations into a predefined number of clusters by repeatedly assigning observations to the nearest centroid and updating the centroid positions."""
)

add_paragraph(
    doc,
    """Educational data mining and learning analytics also demonstrate the usefulness of analysing student records to understand academic patterns. Data preprocessing and appropriate feature selection are important because clustering results depend strongly on the quality and scale of input data."""
)

add_heading(doc, "2.3 Comparative Analysis", 2)

add_paragraph(
    doc,
    """The reviewed technical sources indicate that clustering can be used for numerical pattern discovery, while visualization helps users interpret analytical results. The proposed project combines these concepts with persistent student records and an interactive dashboard."""
)

add_heading(doc, "2.4 Research / Technical Gap", 2)

add_paragraph(
    doc,
    """A simple academic application can be improved by combining similarity-based student grouping with persistent storage, interactive visualization, individual student analysis and report generation. The proposed system addresses this implementation-level gap by integrating these components into one application."""
)

add_heading(doc, "2.5 Chapter Summary", 2)

add_paragraph(
    doc,
    """The literature survey establishes the use of K-Means clustering for grouping numerical data and highlights the importance of preprocessing and visualization. These concepts are applied in the proposed student learning pattern analysis system.""")


# Literature table

add_paragraph(
    doc,
    "Table 2.1 presents selected technical and research references relevant to clustering and visualization."
)

headers = [
    "S.No",
    "Source",
    "Topic",
    "Relevance"
]

rows = [
    [
        "1",
        "MacQueen (1967)",
        "K-Means clustering",
        "Introduced the K-Means approach for grouping observations."
    ],
    [
        "2",
        "Lloyd (1982)",
        "Least-squares quantization",
        "Provides the iterative basis related to K-Means clustering."
    ],
    [
        "3",
        "Jain (2010)",
        "Data clustering",
        "Discusses clustering concepts and limitations."
    ],
    [
        "4",
        "Scikit-learn documentation",
        "K-Means and StandardScaler",
        "Provides implementation support for clustering and feature scaling."
    ],
    [
        "5",
        "Plotly documentation",
        "Interactive visualization",
        "Supports graphical analysis of analytical results."
    ]
]

add_table(doc, headers, rows)


# ============================================================
# CHAPTER 3
# ============================================================

add_heading(doc, "CHAPTER 3", 1)
add_heading(doc, "EXISTING SYSTEM", 1)

add_heading(doc, "3.1 Overview", 2)

add_paragraph(
    doc,
    """In a conventional academic monitoring process, student records are commonly maintained using spreadsheets, documents or separate institutional systems. Teachers inspect marks, attendance and other indicators manually and identify students who may need attention."""
)

add_heading(doc, "3.2 Existing System Architecture", 2)

add_paragraph(
    doc,
    """The general workflow of the existing system is: student records are collected, stored in spreadsheets or databases, manually filtered and analysed, and finally used by teachers to make academic decisions."""
)

add_figure_placeholder(
    doc,
    "3.1",
    "General Workflow of Existing Student Monitoring System"
)

add_heading(doc, "3.3 Working of Existing System", 2)

add_number(doc, "Student information is collected from academic records.")
add_number(doc, "Attendance and marks are maintained.")
add_number(doc, "Teachers inspect individual records.")
add_number(doc, "Students requiring attention are identified manually.")
add_number(doc, "Academic actions are planned based on teacher judgement.")

add_heading(doc, "3.4 Advantages", 2)

for item in [
    "Easy to understand.",
    "Requires minimal technical infrastructure.",
    "Suitable for small datasets.",
    "Teachers can directly interpret student records."
]:
    add_bullet(doc, item)

add_heading(doc, "3.5 Limitations", 2)

for item in [
    "Manual analysis can be time consuming.",
    "Similar student groups may not be immediately visible.",
    "Repeated spreadsheet inspection is required.",
    "There is no automatic clustering.",
    "Visual analysis may be limited.",
    "Report preparation can require additional manual work."
]:
    add_bullet(doc, item)

add_heading(doc, "3.6 Need for Proposed System", 2)

add_paragraph(
    doc,
    """A more organized system can reduce repeated manual analysis by automatically grouping students according to selected numerical features. Interactive graphs can also make the resulting patterns easier to understand.""")


# ============================================================
# CHAPTER 4
# ============================================================

add_heading(doc, "CHAPTER 4", 1)
add_heading(doc, "PROPOSED SYSTEM", 1)

add_heading(doc, "4.1 Overview", 2)

add_paragraph(
    doc,
    """The proposed system is a Student Learning Pattern Analysis application based on K-Means clustering. It accepts student performance data, validates and preprocesses the data, applies clustering and presents the results through an interactive Streamlit dashboard."""
)

add_heading(doc, "4.2 Proposed System Architecture", 2)

add_paragraph(
    doc,
    """The proposed architecture consists of authentication, data access, validation, preprocessing, feature selection, standardization, K-Means clustering, result interpretation, visualization and reporting modules."""
)

add_figure_placeholder(
    doc,
    "4.1",
    "Proposed System Architecture"
)

add_heading(doc, "4.3 Proposed Modules", 2)

modules = [
    "Authentication Module",
    "Student Dashboard",
    "Teacher/Admin Command Center",
    "Data Management Module",
    "Data Preprocessing Module",
    "K-Means Clustering Module",
    "Analytics and Visualization Module",
    "Intervention Management Module",
    "Excel Reporting Module"
]

for module in modules:
    add_bullet(doc, module)

add_heading(doc, "4.4 Working Principle", 2)

add_number(doc, "The user logs into the system according to the assigned role.")
add_number(doc, "The application loads the student dataset.")
add_number(doc, "Input records are validated.")
add_number(doc, "Required numerical features are selected.")
add_number(doc, "The selected features are standardized.")
add_number(doc, "K-Means clustering is applied.")
add_number(doc, "Cluster assignments are analysed.")
add_number(doc, "Results are displayed through interactive graphs.")
add_number(doc, "Teachers can record interventions and feedback.")
add_number(doc, "Student analysis can be exported as an Excel report.")

add_heading(doc, "4.5 Flowchart", 2)

add_paragraph(
    doc,
    """Start → Login → Select Role → Load Data → Validate Data → Preprocess Data → Standardize Features → Apply K-Means → Analyse Clusters → Display Dashboard → Intervention / Report → End"""
)

add_figure_placeholder(
    doc,
    "4.2",
    "Flowchart of Proposed System"
)

add_heading(doc, "4.6 Algorithm", 2)

algorithm_steps = [
    "Select the numerical features required for clustering.",
    "Load the student dataset.",
    "Check the input records for missing or invalid values.",
    "Convert required fields into numerical form.",
    "Standardize the selected features.",
    "Select the number of clusters K.",
    "Initialize K cluster centroids.",
    "Assign each student to the nearest centroid.",
    "Recalculate the centroid of each cluster.",
    "Repeat assignment and centroid calculation until convergence.",
    "Display and interpret the resulting student groups."
]

for i, step in enumerate(algorithm_steps, 1):
    add_paragraph(doc, f"Step {i}: {step}")


add_heading(doc, "4.7 Mathematical Formulation", 2)

add_paragraph(
    doc,
    """For a student vector xᵢ and cluster centroid μⱼ, the Euclidean distance is calculated as:"""
)

add_paragraph(
    doc,
    "d(xᵢ, μⱼ) = √ Σₖ (xᵢₖ − μⱼₖ)²",
    align=WD_ALIGN_PARAGRAPH.CENTER
)

add_paragraph(
    doc,
    """The K-Means objective is to minimize the within-cluster sum of squared distances:"""
)

add_paragraph(
    doc,
    "J = Σᵢ minⱼ ||xᵢ − μⱼ||²",
    align=WD_ALIGN_PARAGRAPH.CENTER
)

add_paragraph(
    doc,
    """Feature standardization is used because different student attributes may be measured on different scales. Standardization helps prevent a feature with a larger numerical range from dominating the distance calculation."""
)

add_heading(doc, "4.8 Advantages of Proposed System", 2)

for item in [
    "Automatically groups similar students.",
    "Provides interactive visualization.",
    "Stores data using SQLite.",
    "Provides role-based access.",
    "Supports individual student analysis.",
    "Allows intervention records.",
    "Provides Excel report generation."
]:
    add_bullet(doc, item)


# ============================================================
# CHAPTER 5
# ============================================================

add_heading(doc, "CHAPTER 5", 1)
add_heading(doc, "SYSTEM REQUIREMENTS", 1)

add_heading(doc, "5.1 Hardware Requirements", 2)

hardware = [
    ["Processor", "Modern dual-core or above processor"],
    ["RAM", "Minimum 4 GB"],
    ["Storage", "At least 5 GB free space"],
    ["Display", "Standard monitor"],
    ["Input", "Keyboard and mouse"],
    ["GPU", "Not required"]
]

add_table(
    doc,
    ["Component", "Requirement"],
    hardware
)

add_heading(doc, "5.2 Software Requirements", 2)

software = [
    ["Operating System", "Windows / Linux"],
    ["Python", "Python 3.x"],
    ["Framework", "Streamlit"],
    ["Database", "SQLite"],
    ["Data Processing", "Pandas, NumPy"],
    ["Machine Learning", "Scikit-learn"],
    ["Visualization", "Plotly"],
    ["Excel Reporting", "openpyxl"]
]

add_table(
    doc,
    ["Software", "Requirement"],
    software
)

add_heading(doc, "5.3 Programming Language", 2)

add_paragraph(
    doc,
    """Python is used as the primary programming language because it provides libraries for data processing, machine learning, visualization, database integration and web application development."""
)

add_heading(doc, "5.4 Development Platform", 2)

add_paragraph(
    doc,
    """The application is developed as a local Streamlit application. SQLite is used for persistent storage and Scikit-learn provides the K-Means clustering implementation.""")


# ============================================================
# CHAPTER 6
# ============================================================

add_heading(doc, "CHAPTER 6", 1)
add_heading(doc, "METHODOLOGY AND IMPLEMENTATION", 1)

add_heading(doc, "6.1 Methodology", 2)

add_paragraph(
    doc,
    """The methodology consists of data collection, validation, preprocessing, feature selection, standardization, clustering, visualization and interpretation. Additional modules provide authentication, database storage, intervention management and report generation."""
)

add_figure_placeholder(
    doc,
    "6.1",
    "Overall Methodology of the Proposed System"
)

add_heading(doc, "6.2 Input / Data Collection", 2)

add_paragraph(
    doc,
    """The application uses structured student records containing academic and learning-related numerical attributes. Data can be maintained within the project database and can also be supplied through CSV-based data management."""
)

add_heading(doc, "6.3 Data Preprocessing", 2)

for item in [
    "Validate required columns.",
    "Check missing and invalid records.",
    "Convert numerical fields into suitable data types.",
    "Select relevant numerical features.",
    "Apply feature standardization."
]:
    add_bullet(doc, item)

add_heading(doc, "6.4 Feature Extraction / Parameter Selection", 2)

add_paragraph(
    doc,
    """Features such as attendance, academic marks, GPA and other available numerical learning indicators can be selected for clustering. Feature selection is important because irrelevant attributes may reduce the usefulness of the resulting groups."""
)

add_heading(doc, "6.5 Module Implementation", 2)

implementation_modules = [
    "Authentication and role management.",
    "Student dashboard.",
    "Teacher and administrator command center.",
    "CSV data validation and management.",
    "Data preprocessing.",
    "K-Means clustering.",
    "Interactive analytics.",
    "Intervention tracking.",
    "Excel report generation."
]

for item in implementation_modules:
    add_bullet(doc, item)

add_heading(doc, "6.6 Hardware Implementation", 2)

add_paragraph(
    doc,
    """No special hardware is required. The system can run on a standard personal computer or laptop with Python installed."""
)

add_heading(doc, "6.7 Software Implementation", 2)

add_paragraph(
    doc,
    """The user interface is developed using Streamlit. Pandas and NumPy are used for data processing. Scikit-learn provides StandardScaler and K-Means clustering. Plotly is used for interactive visualization. SQLite provides persistent storage and openpyxl is used for Excel report generation."""
)

add_heading(doc, "6.8 Implementation Screenshots", 2)

add_paragraph(
    doc,
    """Replace the following placeholders with screenshots captured from the final working application."""
)

add_figure_placeholder(doc, "6.2", "Login Screen")
add_figure_placeholder(doc, "6.3", "Student Dashboard")
add_figure_placeholder(doc, "6.4", "Teacher / Administrator Dashboard")
add_figure_placeholder(doc, "6.5", "K-Means Cluster Visualization")
add_figure_placeholder(doc, "6.6", "Individual Student Analysis")
add_figure_placeholder(doc, "6.7", "Excel Report Generation")

add_heading(doc, "6.9 Testing", 2)

testing_rows = [
    ["TC01", "Login", "Valid credentials", "Dashboard opens", "Pass"],
    ["TC02", "Login", "Invalid credentials", "Error displayed", "Pass"],
    ["TC03", "Data Loading", "Valid student dataset", "Dataset loads", "Pass"],
    ["TC04", "Validation", "Invalid/missing data", "Validation message", "Pass"],
    ["TC05", "Clustering", "Valid numerical features", "Clusters generated", "Pass"],
    ["TC06", "Dashboard", "Clustered dataset", "Graphs displayed", "Pass"],
    ["TC07", "Student Search", "Valid student", "Student analysis displayed", "Pass"],
    ["TC08", "Intervention", "Feedback entry", "Record saved", "Pass"],
    ["TC09", "Excel", "Student report request", "Excel file generated", "Pass"]
]

add_table(
    doc,
    ["Test ID", "Module", "Input", "Expected Result", "Status"],
    testing_rows
)


# ============================================================
# CHAPTER 7
# ============================================================

add_heading(doc, "CHAPTER 7", 1)
add_heading(doc, "RESULTS AND DISCUSSION", 1)

add_heading(doc, "7.1 Experimental Setup", 2)

add_paragraph(
    doc,
    """The application was tested in a local Python environment using the Streamlit framework. Student data was loaded into the application and processed using the preprocessing and clustering pipeline. Standardized numerical features were supplied to the K-Means algorithm."""
)

add_heading(doc, "7.2 Experimental Results", 2)

add_paragraph(
    doc,
    """The system produces cluster assignments for student records. Students with similar values in the selected features are placed into the same cluster. The final cluster counts and numerical performance values should be recorded from the actual dataset used during the final project demonstration."""
)

add_figure_placeholder(
    doc,
    "7.1",
    "Student Cluster Distribution"
)

add_heading(doc, "7.3 Performance Evaluation", 2)

add_paragraph(
    doc,
    """The system can evaluate clustering quality using measures such as the silhouette score. The actual score should be taken from the final experiment and inserted here. The score should be interpreted together with the dataset size, selected features and chosen number of clusters."""
)

add_paragraph(
    doc,
    "Final Experimental Silhouette Score: ____________________"
)

add_heading(doc, "7.4 Graphical Analysis", 2)

add_paragraph(
    doc,
    """The dashboard provides graphical analysis of student performance and cluster patterns. Graphs can be used to compare academic indicators, inspect cluster distribution and observe relationships between selected features."""
)

add_figure_placeholder(
    doc,
    "7.2",
    "Performance Comparison Graph"
)

add_figure_placeholder(
    doc,
    "7.3",
    "Cluster Visualization"
)

add_figure_placeholder(
    doc,
    "7.4",
    "Correlation Analysis"
)

add_heading(doc, "7.5 Comparison with Existing System", 2)

comparison_rows = [
    ["Data Analysis", "Mostly manual", "Automated preprocessing"],
    ["Student Grouping", "Manual interpretation", "K-Means clustering"],
    ["Visualization", "Limited", "Interactive graphs"],
    ["Storage", "Spreadsheet/manual", "SQLite database"],
    ["Intervention Records", "Separate/manual", "Integrated"],
    ["Report Generation", "Manual", "Excel export"]
]

add_table(
    doc,
    ["Feature", "Existing System", "Proposed System"],
    comparison_rows
)

add_heading(doc, "7.6 Discussion", 2)

add_paragraph(
    doc,
    """The results demonstrate that K-Means can be used to organize students into groups based on similarity in selected numerical attributes. The resulting clusters should be treated as data patterns rather than permanent labels of student ability."""
)

add_paragraph(
    doc,
    """The system can support academic monitoring by providing a common analytical view of the dataset. However, teacher judgement remains important when deciding academic interventions."""
)

add_heading(doc, "7.7 Achievement of Objectives", 2)

achievement = [
    ["Objective", "Status"],
    ["Student data collection and storage", "Achieved"],
    ["Data validation and preprocessing", "Achieved"],
    ["K-Means clustering", "Achieved"],
    ["Interactive visualization", "Achieved"],
    ["Individual student analysis", "Achieved"],
    ["Intervention records", "Achieved"],
    ["Excel report generation", "Achieved"]
]

add_table(
    doc,
    ["Objective", "Status"],
    achievement
)


# ============================================================
# CHAPTER 8
# ============================================================

add_heading(doc, "CHAPTER 8", 1)
add_heading(doc, "CONCLUSION AND FUTURE SCOPE", 1)

add_heading(doc, "8.1 Conclusion", 2)

add_paragraph(
    doc,
    """The Student Learning Pattern Analysis Using Clustering project demonstrates a practical application of machine learning for academic data analysis. The system uses K-Means clustering to group students based on selected learning and academic indicators."""
)

add_paragraph(
    doc,
    """The combination of Python, Streamlit, SQLite, Scikit-learn, Pandas and Plotly provides an integrated environment for data processing, clustering, visualization and academic monitoring. The application also supports role-based access, intervention records and Excel report generation."""
)

add_paragraph(
    doc,
    """The system is designed as a decision-support application. Cluster results describe patterns in the available data and should be considered along with teacher judgement and other academic information.""")

add_heading(doc, "8.2 Limitations", 2)

limitations = [
    "The quality of results depends on the quality of the input dataset.",
    "Feature selection can influence the clustering result.",
    "The number of clusters K must be selected appropriately.",
    "Outliers can affect cluster centroids.",
    "Different initialization conditions can influence K-Means results.",
    "Clusters should not be interpreted as permanent measures of student ability.",
    "The early-warning score used by the application is a heuristic indicator and is not a trained probability model."
]

for item in limitations:
    add_bullet(doc, item)

add_heading(doc, "8.3 Future Scope", 2)

future = [
    "Use larger and more diverse student datasets.",
    "Include additional behavioural and learning features.",
    "Automatically determine an appropriate number of clusters.",
    "Compare K-Means with other clustering algorithms.",
    "Develop a supervised early-warning prediction model.",
    "Add time-series analysis of student performance.",
    "Support institutional deployment with stronger privacy and access controls.",
    "Provide more advanced academic recommendation features."
]

for item in future:
    add_bullet(doc, item)


# ============================================================
# REFERENCES
# ============================================================

add_heading(doc, "REFERENCES", 1)

references = [
    "J. MacQueen, “Some Methods for Classification and Analysis of Multivariate Observations,” Proceedings of the Fifth Berkeley Symposium on Mathematical Statistics and Probability, 1967.",
    "S. P. Lloyd, “Least Squares Quantization in PCM,” IEEE Transactions on Information Theory, 1982.",
    "A. K. Jain, “Data Clustering: 50 Years Beyond K-Means,” Pattern Recognition Letters, vol. 31, no. 8, pp. 651–666, 2010.",
    "Scikit-learn Documentation, K-Means Clustering and StandardScaler documentation.",
    "Plotly Documentation, Python Graphing and Interactive Visualization documentation."
]

for ref in references:
    add_paragraph(doc, ref)


# ============================================================
# APPENDIX
# ============================================================

add_heading(doc, "APPENDIX A", 1)
add_heading(doc, "PROJECT SCREENSHOT CHECKLIST", 2)

screenshots = [
    "Login page",
    "Student dashboard",
    "Teacher dashboard",
    "Administrator dashboard",
    "Data management page",
    "Cluster analysis graph",
    "Individual student analysis",
    "Intervention management",
    "Excel report output"
]

for item in screenshots:
    add_bullet(doc, item)

add_heading(doc, "APPENDIX B", 1)
add_heading(doc, "PROJECT TECHNOLOGY STACK", 2)

stack_rows = [
    ["Programming Language", "Python"],
    ["Frontend / UI", "Streamlit"],
    ["Database", "SQLite"],
    ["Data Processing", "Pandas / NumPy"],
    ["Machine Learning", "Scikit-learn"],
    ["Visualization", "Plotly"],
    ["Excel Reporting", "openpyxl"],
    ["Clustering Algorithm", "K-Means"]
]

add_table(
    doc,
    ["Technology", "Usage"],
    stack_rows
)


# ============================================================
# FOOTER
# ============================================================

for section in doc.sections:

    footer = section.footer

    if len(footer.paragraphs) == 0:
        p = footer.add_paragraph()
    else:
        p = footer.paragraphs[0]

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Remove existing footer text
    p.clear()

    run = p.add_run(PROJECT_TITLE)
    set_font(run, 9)


# ============================================================
# GLOBAL FONT
# ============================================================

set_document_font(doc)


# ============================================================
# SAVE
# ============================================================

print()
print("Saving report...")

doc.save(OUTPUT_FILE)

print()
print("=" * 60)
print("REPORT CREATED SUCCESSFULLY")
print("=" * 60)
print()
print(f"File: {OUTPUT_FILE}")
print()
print("Open the generated .docx file in Microsoft Word.")
print("Replace the [ INSERT PROJECT SCREENSHOT / FIGURE HERE ]")
print("sections with screenshots from your actual project.")
print()
